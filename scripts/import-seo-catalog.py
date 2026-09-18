import json
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


CATEGORY_SLUGS = {
    "Potrošni material": "vlozki-in-potrosni-material",
    "Filtri za celo hišo": "filtracija-celega-doma",
    "Podpultni filtri": "filtri-pod-pultom",
    "Testiranje vode": "testiranje-vode",
    "Reverzna osmoza": "reverzna-osmoza",
    "Filtri za pipo": "filtri-za-pipo",
    "Filtri za prho": "filtri-za-prho",
    "Filtri za aparate": "filtri-za-aparate",
    "UV dezinfekcija": "uv-dezinfekcija",
    "Namizni filtri": "namizni-filtri",
    "Mehčalci vode": "mehcalci-vode",
    "Pipe za filtrirano vodo": "pipe-za-filtrirano-vodo",
    "Zaščita in dodatki": "zascita-in-dodatki",
    "Regulacija tlaka": "regulacija-tlaka",
}

FIELD_NAMES = {
    "ID": "id",
    "Prednost": "priority",
    "Kategorija": "sourceCategory",
    "Blagovna znamka": "brand",
    "Naziv izdelka (SL)": "name",
    "Primarna ključna beseda": "primaryKeyword",
    "Sekundarne ključne besede": "secondaryKeywords",
    "Dolgorepe ključne besede": "longTailKeywords",
    "Iskalni namen": "searchIntent",
    "SEO slug": "slug",
    "Predlagani H1": "h1",
    "SEO naslov": "seoTitle",
    "Meta opis": "sourceMetaDescription",
    "Vir": "researchSourceUrl",
    "Opomba": "researchNote",
    "Ocena /100": "seoScore",
}

REQUIRED_SOURCE_FIELDS = tuple(FIELD_NAMES)


def split_keywords(value):
    if not value:
        return []
    return [part.strip() for part in str(value).split(";") if part.strip()]


def main(source_path: Path, output_path: Path) -> None:
    workbook = load_workbook(source_path, read_only=True, data_only=True)
    worksheet = workbook["SEO katalog SL"]
    headers = [cell.value for cell in next(worksheet.iter_rows(min_row=4, max_row=4))]

    missing_headers = [header for header in REQUIRED_SOURCE_FIELDS if header not in headers]
    if missing_headers:
        raise ValueError(f"Missing required columns: {', '.join(missing_headers)}")

    products = []
    for row_index, row in enumerate(worksheet.iter_rows(min_row=5, values_only=True), start=5):
        if not any(value not in (None, "") for value in row):
            continue

        source = dict(zip(headers, row, strict=False))
        missing_values = [
            field for field in REQUIRED_SOURCE_FIELDS if source.get(field) in (None, "")
        ]
        if missing_values:
            raise ValueError(
                f"Row {row_index} is missing: {', '.join(missing_values)}"
            )

        source_category = str(source["Kategorija"]).strip()
        if source_category not in CATEGORY_SLUGS:
            raise ValueError(f"Unmapped category on row {row_index}: {source_category}")

        product = {
            FIELD_NAMES[field]: source[field]
            for field in REQUIRED_SOURCE_FIELDS
            if field not in {"Sekundarne ključne besede", "Dolgorepe ključne besede"}
        }
        product["secondaryKeywords"] = split_keywords(
            source["Sekundarne ključne besede"]
        )
        product["longTailKeywords"] = split_keywords(
            source["Dolgorepe ključne besede"]
        )
        product["categorySlug"] = CATEGORY_SLUGS[source_category]
        product["sourceRow"] = row_index
        products.append(product)

    for key in ("id", "slug"):
        duplicates = [
            value
            for value, count in Counter(product[key] for product in products).items()
            if count > 1
        ]
        if duplicates:
            raise ValueError(f"Duplicate {key} values: {', '.join(duplicates)}")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        json.dumps(products, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )

    print(f"Imported {len(products)} products to {output_path}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: import-seo-catalog.py SOURCE.xlsx OUTPUT.json")

    main(Path(sys.argv[1]), Path(sys.argv[2]))
